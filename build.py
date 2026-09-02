#!/usr/bin/env python3
"""Build Northline_Headcount_Staffing_Forecast.xlsx — HC plan / cost / reforecast portfolio sample."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

yellow = PatternFill("solid", fgColor="FFF2CC")
header_fill = PatternFill("solid", fgColor="1F4E79")
section_fill = PatternFill("solid", fgColor="D6E3F0")
light_gray = PatternFill("solid", fgColor="F5F5F5")
green_fill = PatternFill("solid", fgColor="C6EFCE")
amber_fill = PatternFill("solid", fgColor="FFE699")
red_fill = PatternFill("solid", fgColor="F8CBAD")
tile_fill = PatternFill("solid", fgColor="E9EDF4")
input_font = Font(name="Calibri", size=11, color="0000FF")
black = Font(name="Calibri", size=11, color="000000")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
section_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
bold = Font(name="Calibri", size=11, bold=True)
bold_black = Font(name="Calibri", size=11, bold=True, color="000000")
italic_grey = Font(name="Calibri", size=10, italic=True, color="666666")
small_grey = Font(name="Calibri", size=9, italic=True, color="666666")
link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
money = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
pct = "0.0%"
fte_fmt = "0.0"
num = "#,##0.0"

MONTHS = [
    "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26",
    "Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26",
]
DEPTS = [
    # name, opening FTE, annual loaded $000s, annual attrition
    ("Operations", 120, 72, 0.14),
    ("Supply Chain", 48, 78, 0.12),
    ("Sales", 40, 128, 0.16),
    ("Marketing", 18, 112, 0.11),
    ("Finance", 16, 118, 0.08),
    ("HR", 8, 96, 0.10),
    ("Quality / R&D", 22, 90, 0.09),
    ("G&A / IT", 14, 108, 0.07),
]
N = len(DEPTS)

PLAN_HIRES = [
    [2, 1, 2, 1, 2, 2, 3, 2, 1, 2, 1, 1],
    [1, 0, 1, 1, 0, 1, 1, 1, 0, 1, 1, 0],
    [1, 1, 0, 1, 1, 0, 1, 0, 1, 2, 0, 1],
    [0, 1, 0, 0, 1, 0, 0, 1, 0, 0, 0, 0],
    [0, 0, 1, 0, 0, 0, 0, 1, 0, 0, 0, 0],
    [0, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 1, 0, 0, 1, 0, 1, 0, 0, 0, 0, 0],
    [0, 0, 0, 1, 0, 0, 0, 0, 0, 1, 0, 0],
]
PLAN_OTHER = [
    [0, 0, 1, 0, 0, 0, 0, 1, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
]
# Forecast = YTD actuals through June (as-of = 6) + rest-of-year outlook
FCST_HIRES = [
    [2, 1, 2, 1, 2, 1, 2, 2, 1, 2, 2, 1],
    [1, 0, 1, 0, 1, 1, 1, 1, 0, 1, 1, 0],
    [1, 0, 1, 1, 1, 1, 2, 0, 1, 1, 0, 1],
    [0, 1, 0, 0, 0, 1, 0, 1, 0, 0, 0, 0],
    [0, 0, 1, 0, 0, 0, 0, 0, 0, 1, 0, 0],
    [0, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0, 0],
    [1, 0, 0, 0, 1, 0, 1, 0, 0, 0, 0, 0],
    [0, 0, 0, 1, 0, 0, 0, 0, 1, 0, 0, 0],
]
FCST_OTHER = [
    [0, 1, 0, 0, 0, 1, 0, 1, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
]
OPEN_REQS = [4, 2, 3, 1, 1, 0, 1, 1]

ASSUMP = "01_Assumptions"
PLAN = "02_Headcount_Plan"
COST = "03_Cost_Bridge"
FCST = "04_Reforecast"
DASH = "05_Dashboard"


def style_input(cell):
    cell.fill = yellow
    cell.font = input_font
    cell.border = thin
    cell.alignment = Alignment(horizontal="center")


def style_formula(cell, key=False):
    cell.font = bold_black if key else black
    cell.border = thin
    cell.alignment = Alignment(horizontal="center")
    if key:
        cell.fill = green_fill


def style_header_cell(cell, value):
    cell.value = value
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = thin


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def month_headers(ws, row, start_col=3, include_fy=True, fy_label="FY26"):
    for i, m in enumerate(MONTHS):
        style_header_cell(ws.cell(row, start_col + i), m)
    if include_fy:
        style_header_cell(ws.cell(row, start_col + 12), fy_label)


def landscape(ws, fit_height=0):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.sheet_view.showGridLines = False
    ws.page_setup.horizontalCentered = True
    ws.oddFooter.left.text = "Northline Consumer Products  |  fictional sample"
    ws.oddFooter.right.text = "Page &P of &N"


def col(i):
    return get_column_letter(3 + i)


def shade_section(ws, row, until_col=16):
    ws.cell(row, 2).fill = section_fill
    ws.cell(row, 2).font = section_font
    for c in range(3, until_col):
        ws.cell(row, c).fill = section_fill


wb = Workbook()

# ========== 00_Cover ==========
ws = wb.active
ws.title = "00_Cover"
ws.sheet_properties.tabColor = "1F4E79"
set_col_widths(ws, [4, 92])
landscape(ws, fit_height=1)

ws["B2"] = "Northline Consumer Products"
ws["B2"].font = title_font
ws["B3"] = "Headcount and staffing forecast  ·  plan, loaded cost, reforecast"
ws["B3"].font = section_font
ws["B4"] = "Mid-size CPG  ·  ops + supply chain + commercial + G&A  ·  ~280–310 FTE  ·  FY2026"
ws["B4"].font = italic_grey

ws["B6"] = "What this file is"
ws["B6"].font = bold
ws["B7"] = (
    "A twelve-month headcount plan and reforecast for a fictional CPG company. "
    "It rolls opening FTE + hires − exits by department, applies loaded cost and a mid-year merit, "
    "and compares the original plan to the latest view (YTD actuals through the as-of month, then forecast)."
)
ws["B7"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[7].height = 48

ws["B9"] = "How to use"
ws["B9"].font = bold
ws["B10"] = "1. Edit yellow cells on 01_Assumptions (departments, loaded cost, attrition, raise, as-of month)."
ws["B11"] = "2. Edit plan hires and other exits on 02_Headcount_Plan. Opening, attrition exits, and ending FTE are formulas."
ws["B12"] = "3. Read the loaded-cost rollforward and plan vs forecast $ on 03_Cost_Bridge."
ws["B13"] = "4. Edit forecast hires / other exits on 04_Reforecast; compare latest view vs original plan."
ws["B14"] = "5. Use 05_Dashboard as the one-pager."

ws["B16"] = "File conventions"
ws["B16"].font = bold
ws["B17"] = "Yellow cells with blue font = inputs. Black font = formulas. Green cells = key outputs."
ws["B17"].fill = yellow
ws["B17"].font = Font(name="Calibri", size=11, color="0000FF", bold=True)
ws["B18"] = "All figures are fictional. There is no employer data in this file. Cost in $000s. FTE in heads (one decimal)."

ws["B20"] = "Portfolio"
ws["B20"].font = bold
ws["B21"] = "Sai Siri Bandaru — Financial Analyst | FP&A | forecasting, variance analysis, Excel"
ws["B22"] = "https://github.com/saisiri-bandaru"
ws["B22"].font = link_font

# ========== 01_Assumptions ==========
ws = wb.create_sheet(ASSUMP)
ws.sheet_properties.tabColor = "F7C948"
set_col_widths(ws, [4, 22, 16, 24, 18, 18, 16, 16, 14])
landscape(ws)

ws["B2"] = "Assumptions"
ws["B2"].font = title_font
ws["B3"] = "Yellow + blue = inputs. Change these; plan, cost, reforecast, and dashboard recalculate."
ws["B3"].font = italic_grey

ws["B5"] = "As-of / toggle month (1 = Jan … 12 = Dec)"
ws["C5"] = 6
style_input(ws["C5"]); ws["C5"].number_format = "0"
ws["D5"] = '=INDEX({"Jan-26","Feb-26","Mar-26","Apr-26","May-26","Jun-26","Jul-26","Aug-26","Sep-26","Oct-26","Nov-26","Dec-26"},1,C5)'
style_formula(ws["D5"])
ws["E5"] = "Months ≤ as-of on the reforecast are YTD actuals; months after are forecast."
ws["E5"].font = small_grey

ws["B6"] = "Merit / raise %"
ws["C6"] = 0.035
style_input(ws["C6"]); ws["C6"].number_format = pct
ws["B7"] = "Raise effective month (1–12)"
ws["C7"] = 4
style_input(ws["C7"]); ws["C7"].number_format = "0"
ws["D7"] = "April — loaded cost steps up from this month."
ws["D7"].font = small_grey
ws["B8"] = "Hire lead time (months)"
ws["C8"] = 2
style_input(ws["C8"]); ws["C8"].number_format = "0"
ws["D8"] = "Hires are recorded in the start month (already net of lead time). Open reqs are the pipeline."
ws["D8"].font = small_grey

ws["B10"] = "Department drivers"
ws["B10"].font = section_font
ws["B10"].fill = section_fill
for c in range(3, 8):
    ws.cell(10, c).fill = section_fill

headers = ["Department", "Opening FTE", "Annual loaded cost ($000s)", "Annual attrition %", "Monthly attrition %", "Open reqs (pipeline)"]
for i, h in enumerate(headers):
    style_header_cell(ws.cell(11, 2 + i), h)

for i, (name, opening, loaded, attr) in enumerate(DEPTS):
    r = 12 + i
    ws.cell(r, 2, name).font = bold
    ws.cell(r, 2).border = thin
    c = ws.cell(r, 3, opening); style_input(c); c.number_format = fte_fmt
    c = ws.cell(r, 4, loaded); style_input(c); c.number_format = money
    c = ws.cell(r, 5, attr); style_input(c); c.number_format = pct
    c = ws.cell(r, 6, f"=E{r}/12"); style_formula(c); c.number_format = pct
    c = ws.cell(r, 7, OPEN_REQS[i]); style_input(c); c.number_format = "0"

ws["B20"] = "Total"
ws["B20"].font = bold
ws["C20"] = "=SUM(C12:C19)"; style_formula(ws["C20"], key=True); ws["C20"].number_format = fte_fmt
ws["D20"] = "=SUMPRODUCT(C12:C19,D12:D19)/C20"; style_formula(ws["D20"]); ws["D20"].number_format = money
ws["E20"] = "=SUMPRODUCT(C12:C19,E12:E19)/C20"; style_formula(ws["E20"]); ws["E20"].number_format = pct
ws["F20"] = "=E20/12"; style_formula(ws["F20"]); ws["F20"].number_format = pct
ws["G20"] = "=SUM(G12:G19)"; style_formula(ws["G20"], key=True); ws["G20"].number_format = "0"

ws["B22"] = "Loaded cost is fully loaded (salary + benefits + taxes + typical bonus). Plant roles sit below corporate / sales."
ws["B22"].font = small_grey
ws["B23"] = "Monthly attrition used in the plan = annual attrition ÷ 12 (straight-line). Ending FTE is allowed to be fractional for planning."
ws["B23"].font = small_grey
ws["B24"] = "Company-wide average loaded cost (opening mix)"
ws["C24"] = "=D20"
style_formula(ws["C24"]); ws["C24"].number_format = money

# ========== 02_Headcount_Plan ==========
ws = wb.create_sheet(PLAN)
ws.sheet_properties.tabColor = "5B9BD5"
set_col_widths(ws, [4, 28] + [11] * 13)
landscape(ws)
ws.freeze_panes = "C8"

ws["B2"] = "Headcount plan (original)"
ws["B2"].font = title_font
ws["B3"] = "Hires and other exits are inputs. Opening FTE, attrition exits, total exits, and ending FTE are formulas."
ws["B3"].font = italic_grey

# Block helper
def write_hc_block(ws, title_row, title, row_labels, kind, src_matrix=None, src_sheet=None):
    """kind: 'open', 'hires', 'attrition', 'other', 'exits', 'end' """
    ws.cell(title_row, 2, title)
    shade_section(ws, title_row)
    month_headers(ws, title_row + 1)
    first = title_row + 2
    for d, (name, *_rest) in enumerate(DEPTS):
        r = first + d
        ws.cell(r, 2, name).border = thin
        for m in range(12):
            cl = col(m)
            p = col(m - 1) if m else None
            cell = ws.cell(r, 3 + m)
            ar = 12 + d  # assumptions dept row
            if kind == "open":
                if m == 0:
                    cell.value = f"='{ASSUMP}'!C{ar}"
                else:
                    # prior ending is 50 rows below opening? we'll pass end_row_offset via src_sheet as int
                    cell.value = f"={p}{src_matrix}"  # src_matrix is ending row for this dept
                style_formula(cell)
            elif kind == "hires":
                cell.value = src_matrix[d][m]
                style_input(cell)
            elif kind == "attrition":
                # ROUND(opening * monthly attrition, 1)
                cell.value = f"=ROUND({cl}{src_matrix}*'{ASSUMP}'!$F${ar},1)"
                style_formula(cell)
            elif kind == "other":
                cell.value = src_matrix[d][m]
                style_input(cell)
            elif kind == "exits":
                # attrition row + other row
                cell.value = f"={cl}{src_matrix[0]}+{cl}{src_matrix[1]}"
                style_formula(cell)
            elif kind == "end":
                # open + hires - exits
                cell.value = f"={cl}{src_matrix[0]}+{cl}{src_matrix[1]}-{cl}{src_matrix[2]}"
                style_formula(cell, key=True)
            cell.number_format = fte_fmt
        # FY: average ending / sum hires / etc
        if kind in ("hires", "other", "attrition", "exits"):
            cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
            style_formula(cell, key=(kind == "exits"))
        elif kind == "open":
            cell = ws.cell(r, 15, f"=C{r}")
            style_formula(cell)
        else:
            cell = ws.cell(r, 15, f"=N{r}")
            style_formula(cell, key=True)
        cell.number_format = fte_fmt
    tot = first + N
    ws.cell(tot, 2, "Total").font = bold
    ws.cell(tot, 2).border = thin
    for m in range(13):
        cl = get_column_letter(3 + m)
        cell = ws.cell(tot, 3 + m, f"=SUM({cl}{first}:{cl}{first+N-1})")
        style_formula(cell, key=True)
        cell.number_format = fte_fmt
    return first, tot


# Layout:
# Opening rows 8-16 (8 depts + we put title at 6, headers 7, depts 8-15, total 16)
# But opening needs to reference ending rows which we know in advance.
# Opening block title row 6, headers 7, depts 8-15, total 16
# Hires title 18, headers 19, depts 20-27, total 28
# Attrition title 30, headers 31, depts 32-39, total 40
# Other title 42, headers 43, depts 44-51, total 52
# Total exits title 54, headers 55, depts 56-63, total 64
# Ending title 66, headers 67, depts 68-75, total 76

# Opening FTE — Jan from assumptions, later months from prior ending (row 68+d)
ws.cell(6, 2, "Opening FTE")
shade_section(ws, 6)
month_headers(ws, 7)
for d, (name, *_) in enumerate(DEPTS):
    r = 8 + d
    end_r = 68 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m)
        if m == 0:
            cell.value = f"='{ASSUMP}'!C{12+d}"
        else:
            cell.value = f"={col(m-1)}{end_r}"
        style_formula(cell)
        cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=C{r}")
    style_formula(cell); cell.number_format = fte_fmt
ws["B16"] = "Total"; ws["B16"].font = bold; ws["B16"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(16, 3 + m, f"=SUM({cl}8:{cl}15)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(18, 2, "Hires (starts) — input")
shade_section(ws, 18)
month_headers(ws, 19)
for d, (name, *_) in enumerate(DEPTS):
    r = 20 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m, PLAN_HIRES[d][m])
        style_input(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B28"] = "Total"; ws["B28"].font = bold; ws["B28"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(28, 3 + m, f"=SUM({cl}20:{cl}27)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(30, 2, "Modeled attrition exits  = ROUND(opening × monthly attrition, 1)")
shade_section(ws, 30)
month_headers(ws, 31)
for d, (name, *_) in enumerate(DEPTS):
    r = 32 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"=ROUND({cl}{8+d}*'01_Assumptions'!$F${12+d},1)")
        style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B40"] = "Total"; ws["B40"].font = bold; ws["B40"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(40, 3 + m, f"=SUM({cl}32:{cl}39)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(42, 2, "Other exits (restructuring / transfers-out) — input")
shade_section(ws, 42)
month_headers(ws, 43)
for d, (name, *_) in enumerate(DEPTS):
    r = 44 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m, PLAN_OTHER[d][m])
        style_input(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B52"] = "Total"; ws["B52"].font = bold; ws["B52"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(52, 3 + m, f"=SUM({cl}44:{cl}51)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(54, 2, "Total exits  = attrition + other")
shade_section(ws, 54)
month_headers(ws, 55)
for d, (name, *_) in enumerate(DEPTS):
    r = 56 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"={cl}{32+d}+{cl}{44+d}")
        style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B64"] = "Total"; ws["B64"].font = bold; ws["B64"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(64, 3 + m, f"=SUM({cl}56:{cl}63)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(66, 2, "Ending FTE  = opening + hires − total exits")
shade_section(ws, 66)
month_headers(ws, 67)
for d, (name, *_) in enumerate(DEPTS):
    r = 68 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"={cl}{8+d}+{cl}{20+d}-{cl}{56+d}")
        style_formula(cell, key=True); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=N{r}")
    style_formula(cell, key=True); cell.number_format = fte_fmt
ws["B76"] = "Total"; ws["B76"].font = bold; ws["B76"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(76, 3 + m, f"=SUM({cl}68:{cl}75)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws["B78"] = "Net change (ending − opening)"
shade_section(ws, 78)
month_headers(ws, 79)
ws["B80"] = "Total net hires / (exits)"
for m in range(12):
    cl = col(m)
    cell = ws.cell(80, 3 + m, f"={cl}76-{cl}16")
    style_formula(cell); cell.number_format = fte_fmt
ws["O80"] = "=N76-C16"
style_formula(ws["O80"], key=True); ws["O80"].number_format = fte_fmt

ws["B82"] = "Hires are starts, not requisitions. A 2-month lead time is already baked into the month a hire is booked."
ws["B82"].font = small_grey

# ========== 03_Cost_Bridge ==========
ws = wb.create_sheet(COST)
ws.sheet_properties.tabColor = "70AD47"
set_col_widths(ws, [4, 32] + [12] * 13)
landscape(ws)
ws.freeze_panes = "C8"

ws["B2"] = "Loaded cost rollforward  ($000s)"
ws["B2"].font = title_font
ws["B3"] = "Monthly cost = average FTE × (annual loaded / 12) × raise factor. Raise factor steps up in the merit month."
ws["B3"].font = italic_grey

ws.cell(5, 2, "Raise factor by month")
shade_section(ws, 5)
month_headers(ws, 6)
ws["B7"] = "Raise factor"
for m in range(12):
    # month index m+1
    cell = ws.cell(7, 3 + m, f'=IF({m+1}>=\'{ASSUMP}\'!$C$7,1+\'{ASSUMP}\'!$C$6,1)')
    style_formula(cell); cell.number_format = "0.000"
ws["O7"] = "=N7"
style_formula(ws["O7"]); ws["O7"].number_format = "0.000"

ws.cell(9, 2, "Plan loaded cost by department ($000s)")
shade_section(ws, 9)
month_headers(ws, 10, fy_label="FY26 $")
for d, (name, *_) in enumerate(DEPTS):
    r = 11 + d
    ws.cell(r, 2, name).border = thin
    ar = 12 + d
    open_r = 8 + d
    end_r = 68 + d
    for m in range(12):
        cl = col(m)
        # avg FTE * annual/12 * raise factor
        f = f"=(('{PLAN}'!{cl}{open_r}+'{PLAN}'!{cl}{end_r})/2)*('{ASSUMP}'!$D${ar}/12)*{cl}$7"
        cell = ws.cell(r, 3 + m, f)
        style_formula(cell); cell.number_format = money
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell, key=True); cell.number_format = money
ws["B19"] = "Plan total"
ws["B19"].font = bold; ws["B19"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(19, 3 + m, f"=SUM({cl}11:{cl}18)")
    style_formula(cell, key=True); cell.number_format = money

ws.cell(21, 2, "Forecast loaded cost by department ($000s)  — FTE from 04_Reforecast")
shade_section(ws, 21)
month_headers(ws, 22, fy_label="FY26 $")
for d, (name, *_) in enumerate(DEPTS):
    r = 23 + d
    ws.cell(r, 2, name).border = thin
    ar = 12 + d
    # reforecast opening rows 8-15, ending 68-75 (same layout as plan)
    open_r = 8 + d
    end_r = 68 + d
    for m in range(12):
        cl = col(m)
        f = f"=(('{FCST}'!{cl}{open_r}+'{FCST}'!{cl}{end_r})/2)*('{ASSUMP}'!$D${ar}/12)*{cl}$7"
        cell = ws.cell(r, 3 + m, f)
        style_formula(cell); cell.number_format = money
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell, key=True); cell.number_format = money
ws["B31"] = "Forecast total"
ws["B31"].font = bold; ws["B31"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(31, 3 + m, f"=SUM({cl}23:{cl}30)")
    style_formula(cell, key=True); cell.number_format = money

ws.cell(33, 2, "Variance  (forecast − plan)  $000s")
shade_section(ws, 33)
month_headers(ws, 34, fy_label="FY26 $")
for d, (name, *_) in enumerate(DEPTS):
    r = 35 + d
    ws.cell(r, 2, name).border = thin
    plan_r = 11 + d
    fcst_r = 23 + d
    for m in range(13):
        cl = get_column_letter(3 + m)
        cell = ws.cell(r, 3 + m, f"={cl}{fcst_r}-{cl}{plan_r}")
        style_formula(cell); cell.number_format = money
ws["B43"] = "Variance total"
ws["B43"].font = bold; ws["B43"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(43, 3 + m, f"={cl}31-{cl}19")
    style_formula(cell, key=True); cell.number_format = money

ws.conditional_formatting.add("C35:O43", FormulaRule(formula=["C35>5"], fill=red_fill))
ws.conditional_formatting.add("C35:O43", FormulaRule(formula=["C35<-5"], fill=green_fill))

ws.cell(45, 2, "FY cost bridge (plan) — illustrative, $000s")
shade_section(ws, 45)
style_header_cell(ws.cell(46, 2), "Bridge step")
style_header_cell(ws.cell(46, 3), "$000s")
ws["B47"] = "Opening run-rate  (opening FTE × pre-raise monthly cost × 12)"
ws["C47"] = f"=SUMPRODUCT('{ASSUMP}'!C12:C19,'{ASSUMP}'!D12:D19)"
style_formula(ws["C47"]); ws["C47"].number_format = money
ws["B48"] = "Merit on opening base  (raise × months-in-force / 12)"
ws["C48"] = f"=C47*'{ASSUMP}'!C6*(13-'{ASSUMP}'!C7)/12"
style_formula(ws["C48"]); ws["C48"].number_format = money
ws["B49"] = "In-year volume / mix / timing  (plug to plan FY cost)"
ws["C49"] = "=O19-C47-C48"
style_formula(ws["C49"]); ws["C49"].number_format = money
ws["B50"] = "Plan FY loaded cost (check)"
ws["C50"] = "=C47+C48+C49"
style_formula(ws["C50"], key=True); ws["C50"].number_format = money
ws["D50"] = '=IF(ABS(C50-O19)<1,"Ties to plan FY","Check")'
style_formula(ws["D50"])

ws["B52"] = "Average FTE timing means a hire mid-month (or an exit) costs half a month — standard FP&A convention."
ws["B52"].font = small_grey

# ========== 04_Reforecast ==========
ws = wb.create_sheet(FCST)
ws.sheet_properties.tabColor = "ED7D31"
set_col_widths(ws, [4, 32] + [11] * 13)
landscape(ws)
ws.freeze_panes = "C8"

ws["B2"] = "Reforecast — latest view vs original plan"
ws["B2"].font = title_font
ws["B3"] = "Same mechanics as the plan. Hires and other exits are the latest view (YTD actuals through the as-of month, then forecast)."
ws["B3"].font = italic_grey
ws["B4"] = 'As-of month'
ws["C4"] = f"='{ASSUMP}'!D5"
style_formula(ws["C4"], key=True)

# Month status row
ws.cell(6, 2, "Month status  (Actual if month ≤ as-of, else Forecast)")
shade_section(ws, 6)
month_headers(ws, 7, include_fy=False)
ws["B8"] = "Status"
for m in range(12):
    cell = ws.cell(8, 3 + m, f'=IF({m+1}<=\'{ASSUMP}\'!$C$5,"Actual","Forecast")')
    style_formula(cell)

# Opening 10-19 (title 10, hdr 11, depts 12-19? Keep SAME row numbers as plan for cost formulas: open 8-15, hires 20-27, attr 32-39, other 44-51, exits 56-63, end 68-75
# Cost bridge references FCST open 8+d and end 68+d. So I MUST use the same row numbers as the plan sheet.
# That means I should not have used rows 6-8 for status... Plan has opening at row 8.
# I'll put status on row 5 (headers already at 7 for plan). Let me restructure reforecast to MATCH plan rows exactly.

# Wait - I already wrote rows 2-8. Plan opening is row 8. Conflict.
# Cost uses '{FCST}'!{cl}{open_r} with open_r = 8+d. I need opening depts at rows 8-15 on FCST too.
# Let me move status to columns Q/R or put it at row 80+.
# I'll clear row 6-8 usage and put status at the bottom. Rebuild opening at 8-16 like plan.

# Overwrite: put a thin status line in row 5 using D4 already.
for m in range(12):
    cell = ws.cell(5, 3 + m, f'=IF({m+1}<=\'{ASSUMP}\'!$C$5,"Actual","Forecast")')
    style_formula(cell)
    cell.font = Font(name="Calibri", size=9, italic=True, color="1F4E79")
style_header_cell = style_header_cell  # keep

# Clear the earlier section at 6-8 by rewriting plan-identical blocks starting at 6.

ws.cell(6, 2, "Opening FTE (latest view)")
shade_section(ws, 6)
month_headers(ws, 7)
for d, (name, *_) in enumerate(DEPTS):
    r = 8 + d
    end_r = 68 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m)
        if m == 0:
            cell.value = f"='{ASSUMP}'!C{12+d}"
        else:
            cell.value = f"={col(m-1)}{end_r}"
        style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=C{r}")
    style_formula(cell); cell.number_format = fte_fmt
ws["B16"] = "Total"; ws["B16"].font = bold; ws["B16"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(16, 3 + m, f"=SUM({cl}8:{cl}15)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(18, 2, "Hires (starts) — latest view, input")
shade_section(ws, 18)
month_headers(ws, 19)
for d, (name, *_) in enumerate(DEPTS):
    r = 20 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m, FCST_HIRES[d][m])
        style_input(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B28"] = "Total"; ws["B28"].font = bold; ws["B28"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(28, 3 + m, f"=SUM({cl}20:{cl}27)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(30, 2, "Modeled attrition exits  = ROUND(opening × monthly attrition, 1)")
shade_section(ws, 30)
month_headers(ws, 31)
for d, (name, *_) in enumerate(DEPTS):
    r = 32 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"=ROUND({cl}{8+d}*'01_Assumptions'!$F${12+d},1)")
        style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B40"] = "Total"; ws["B40"].font = bold; ws["B40"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(40, 3 + m, f"=SUM({cl}32:{cl}39)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(42, 2, "Other exits — latest view, input")
shade_section(ws, 42)
month_headers(ws, 43)
for d, (name, *_) in enumerate(DEPTS):
    r = 44 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cell = ws.cell(r, 3 + m, FCST_OTHER[d][m])
        style_input(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B52"] = "Total"; ws["B52"].font = bold; ws["B52"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(52, 3 + m, f"=SUM({cl}44:{cl}51)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(54, 2, "Total exits  = attrition + other")
shade_section(ws, 54)
month_headers(ws, 55)
for d, (name, *_) in enumerate(DEPTS):
    r = 56 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"={cl}{32+d}+{cl}{44+d}")
        style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=SUM(C{r}:N{r})")
    style_formula(cell); cell.number_format = fte_fmt
ws["B64"] = "Total"; ws["B64"].font = bold; ws["B64"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(64, 3 + m, f"=SUM({cl}56:{cl}63)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(66, 2, "Ending FTE (latest view)  = opening + hires − total exits")
shade_section(ws, 66)
month_headers(ws, 67)
for d, (name, *_) in enumerate(DEPTS):
    r = 68 + d
    ws.cell(r, 2, name).border = thin
    for m in range(12):
        cl = col(m)
        cell = ws.cell(r, 3 + m, f"={cl}{8+d}+{cl}{20+d}-{cl}{56+d}")
        style_formula(cell, key=True); cell.number_format = fte_fmt
    cell = ws.cell(r, 15, f"=N{r}")
    style_formula(cell, key=True); cell.number_format = fte_fmt
ws["B76"] = "Total"; ws["B76"].font = bold; ws["B76"].border = thin
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(76, 3 + m, f"=SUM({cl}68:{cl}75)")
    style_formula(cell, key=True); cell.number_format = fte_fmt

# Comparison vs plan
ws.cell(78, 2, "Ending FTE variance vs plan  (latest − plan)")
shade_section(ws, 78)
month_headers(ws, 79)
ws["B80"] = "Total FTE variance"
for m in range(13):
    cl = get_column_letter(3 + m)
    cell = ws.cell(80, 3 + m, f"={cl}76-'{PLAN}'!{cl}76")
    style_formula(cell, key=True); cell.number_format = fte_fmt

ws.cell(82, 2, "Year-end FTE by department — plan vs latest")
shade_section(ws, 82)
style_header_cell(ws.cell(83, 2), "Department")
style_header_cell(ws.cell(83, 3), "Plan YE")
style_header_cell(ws.cell(83, 4), "Latest YE")
style_header_cell(ws.cell(83, 5), "Var FTE")
style_header_cell(ws.cell(83, 6), "Plan FY $")
style_header_cell(ws.cell(83, 7), "Latest FY $")
style_header_cell(ws.cell(83, 8), "Var $")
for d, (name, *_) in enumerate(DEPTS):
    r = 84 + d
    ws.cell(r, 2, name).border = thin
    cell = ws.cell(r, 3, f"='{PLAN}'!N{68+d}"); style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 4, f"=N{68+d}"); style_formula(cell, key=True); cell.number_format = fte_fmt
    cell = ws.cell(r, 5, f"=D{r}-C{r}"); style_formula(cell); cell.number_format = fte_fmt
    cell = ws.cell(r, 6, f"='{COST}'!O{11+d}"); style_formula(cell); cell.number_format = money
    cell = ws.cell(r, 7, f"='{COST}'!O{23+d}"); style_formula(cell); cell.number_format = money
    cell = ws.cell(r, 8, f"=G{r}-F{r}"); style_formula(cell); cell.number_format = money
ws["B92"] = "Total"; ws["B92"].font = bold; ws["B92"].border = thin
ws["C92"] = f"='{PLAN}'!N76"; style_formula(ws["C92"], key=True); ws["C92"].number_format = fte_fmt
ws["D92"] = "=N76"; style_formula(ws["D92"], key=True); ws["D92"].number_format = fte_fmt
ws["E92"] = "=D92-C92"; style_formula(ws["E92"], key=True); ws["E92"].number_format = fte_fmt
ws["F92"] = f"='{COST}'!O19"; style_formula(ws["F92"]); ws["F92"].number_format = money
ws["G92"] = f"='{COST}'!O31"; style_formula(ws["G92"]); ws["G92"].number_format = money
ws["H92"] = "=G92-F92"; style_formula(ws["H92"], key=True); ws["H92"].number_format = money

ws["B94"] = "Positive FTE variance = running above the original plan. Positive $ variance = cost above plan."
ws["B94"].font = small_grey

# ========== 05_Dashboard ==========
ws = wb.create_sheet(DASH)
ws.sheet_properties.tabColor = "1F4E79"
set_col_widths(ws, [4, 22, 16, 3, 24, 16, 3, 22, 16, 3, 20, 14] + [11] * 8)
landscape(ws, fit_height=1)

ws["B2"] = "Headcount dashboard"
ws["B2"].font = title_font
ws["B3"] = "As-of month"
ws["C3"] = f"='{ASSUMP}'!D5"
style_formula(ws["C3"], key=True)
ws["E3"] = "Yellow + blue = inputs (on other tabs). This page is formulas."
ws["E3"].font = small_grey

def idx_plan(row):
    return f"=INDEX('{PLAN}'!C{row}:N{row},1,'{ASSUMP}'!$C$5)"

def idx_fcst(row):
    return f"=INDEX('{FCST}'!C{row}:N{row},1,'{ASSUMP}'!$C$5)"

def idx_cost(row):
    return f"=INDEX('{COST}'!C{row}:N{row},1,'{ASSUMP}'!$C$5)"

# Tiles
ws["B5"] = "Ending FTE (latest)"
ws["B5"].font = header_font; ws["B5"].fill = header_fill; ws["C5"].fill = header_fill
ws["B6"] = "Actual / forecast"; ws["B6"].fill = tile_fill
ws["C6"] = idx_fcst(76); style_formula(ws["C6"], key=True); ws["C6"].number_format = fte_fmt
ws["B7"] = "Original plan"; ws["B7"].fill = tile_fill
ws["C7"] = idx_plan(76); style_formula(ws["C7"]); ws["C7"].number_format = fte_fmt
ws["B8"] = "Var vs plan"; ws["B8"].fill = tile_fill
ws["C8"] = "=C6-C7"; style_formula(ws["C8"]); ws["C8"].number_format = fte_fmt

ws["E5"] = "FY loaded cost ($000s)"
ws["E5"].font = header_font; ws["E5"].fill = header_fill; ws["F5"].fill = header_fill
ws["E6"] = "Forecast FY $"; ws["E6"].fill = tile_fill
ws["F6"] = f"='{COST}'!O31"; style_formula(ws["F6"], key=True); ws["F6"].number_format = money
ws["E7"] = "Plan FY $"; ws["E7"].fill = tile_fill
ws["F7"] = f"='{COST}'!O19"; style_formula(ws["F7"]); ws["F7"].number_format = money
ws["E8"] = "Var $"; ws["E8"].fill = tile_fill
ws["F8"] = "=F6-F7"; style_formula(ws["F8"]); ws["F8"].number_format = money

ws["H5"] = "YTD net / pipeline"
ws["H5"].font = header_font; ws["H5"].fill = header_fill; ws["I5"].fill = header_fill
ws["H6"] = "YTD hires (latest)"; ws["H6"].fill = tile_fill
ws["I6"] = f"=SUMIF(C16:N16,\"Actual\",'{FCST}'!C28:N28)"
# status is on row 5 of FCST C5:N5
ws["I6"] = f"=SUMIF('{FCST}'!C5:N5,\"Actual\",'{FCST}'!C28:N28)"
style_formula(ws["I6"], key=True); ws["I6"].number_format = fte_fmt
ws["H7"] = "YTD exits (latest)"; ws["H7"].fill = tile_fill
ws["I7"] = f"=SUMIF('{FCST}'!C5:N5,\"Actual\",'{FCST}'!C64:N64)"
style_formula(ws["I7"]); ws["I7"].number_format = fte_fmt
ws["H8"] = "Open reqs"; ws["H8"].fill = tile_fill
ws["I8"] = f"='{ASSUMP}'!G20"; style_formula(ws["I8"]); ws["I8"].number_format = "0"

ws["K5"] = "Attrition (annualized)"
ws["K5"].font = header_font; ws["K5"].fill = header_fill; ws["L5"].fill = header_fill
ws["K6"] = "As-of month"; ws["K6"].fill = tile_fill
ws["L6"] = f"=IF({idx_fcst(16)[1:]}=0,0,{idx_fcst(64)}/{idx_fcst(16)}*12)"
# idx_fcst returns formula starting with = so nested IF is messy. Write explicitly.
ws["L6"] = f"=IF(INDEX('{FCST}'!C16:N16,1,'{ASSUMP}'!$C$5)=0,0,INDEX('{FCST}'!C64:N64,1,'{ASSUMP}'!$C$5)/INDEX('{FCST}'!C16:N16,1,'{ASSUMP}'!$C$5)*12)"
style_formula(ws["L6"], key=True); ws["L6"].number_format = pct
ws["K7"] = "Company target"; ws["K7"].fill = tile_fill
ws["L7"] = f"='{ASSUMP}'!E20"; style_formula(ws["L7"]); ws["L7"].number_format = pct
ws["K8"] = "Hire lead time (mo)"; ws["K8"].fill = tile_fill
ws["L8"] = f"='{ASSUMP}'!C8"; style_formula(ws["L8"]); ws["L8"].number_format = "0"

# Trend data
ws["B10"] = "Monthly totals (feeds charts)"
ws["B10"].font = section_font
ws["B10"].fill = section_fill
for c in range(3, 16):
    ws.cell(10, c).fill = section_fill
style_header_cell(ws.cell(11, 2), "Metric")
for i, m in enumerate(MONTHS):
    style_header_cell(ws.cell(11, 3 + i), m)

ws["B12"] = "Plan ending FTE"
ws["B13"] = "Latest ending FTE"
ws["B14"] = "Plan monthly $"
ws["B15"] = "Latest monthly $"
ws["B16"] = "Status"
for m in range(12):
    cl = col(m)
    c = 3 + m
    ws.cell(12, c, f"='{PLAN}'!{cl}76"); style_formula(ws.cell(12, c)); ws.cell(12, c).number_format = fte_fmt
    ws.cell(13, c, f"='{FCST}'!{cl}76"); style_formula(ws.cell(13, c)); ws.cell(13, c).number_format = fte_fmt
    ws.cell(14, c, f"='{COST}'!{cl}19"); style_formula(ws.cell(14, c)); ws.cell(14, c).number_format = money
    ws.cell(15, c, f"='{COST}'!{cl}31"); style_formula(ws.cell(15, c)); ws.cell(15, c).number_format = money
    ws.cell(16, c, f"='{FCST}'!{cl}5"); style_formula(ws.cell(16, c))

chart1 = LineChart()
chart1.title = "Ending FTE — plan vs latest"
chart1.style = 10
chart1.y_axis.title = "FTE"
chart1.height = 7
chart1.width = 14
chart1.legend.position = "b"
data = Reference(ws, min_col=2, min_row=12, max_col=14, max_row=13)
cats = Reference(ws, min_col=3, min_row=11, max_col=14)
chart1.add_data(data, from_rows=True, titles_from_data=True)
chart1.set_categories(cats)
ws.add_chart(chart1, "B18")

chart2 = LineChart()
chart2.title = "Monthly loaded cost — plan vs latest ($000s)"
chart2.style = 12
chart2.y_axis.title = "$000s"
chart2.height = 7
chart2.width = 14
chart2.legend.position = "b"
data2 = Reference(ws, min_col=2, min_row=14, max_col=14, max_row=15)
chart2.add_data(data2, from_rows=True, titles_from_data=True)
chart2.set_categories(cats)
ws.add_chart(chart2, "H18")

# Dept bar at YE
ws["B34"] = "Year-end FTE by department"
ws["B34"].font = section_font
style_header_cell(ws.cell(35, 2), "Department")
style_header_cell(ws.cell(35, 3), "Plan")
style_header_cell(ws.cell(35, 4), "Latest")
for d, (name, *_) in enumerate(DEPTS):
    r = 36 + d
    ws.cell(r, 2, name).border = thin
    cell = ws.cell(r, 3, f"='{FCST}'!C{84+d}"); style_formula(cell); cell.number_format = fte_fmt
    # wait FCST C84 is plan YE already. Use C84 and D84
    cell.value = f"='{FCST}'!C{84+d}"
    cell = ws.cell(r, 4, f"='{FCST}'!D{84+d}"); style_formula(cell); cell.number_format = fte_fmt

bar = BarChart()
bar.type = "col"
bar.grouping = "clustered"
bar.title = "Year-end FTE by department"
bar.style = 10
bar.y_axis.title = "FTE"
bar.height = 7
bar.width = 14
bar.legend.position = "b"
data3 = Reference(ws, min_col=3, min_row=35, max_col=4, max_row=43)
cats3 = Reference(ws, min_col=2, min_row=36, max_row=43)
bar.add_data(data3, titles_from_data=True)
bar.set_categories(cats3)
ws.add_chart(bar, "F34")

ws["B46"] = "Toggle the as-of month on Assumptions to walk the latest view from YTD actuals into the remaining-year forecast."
ws["B46"].font = small_grey

# ========== 06_Data_Dictionary ==========
ws = wb.create_sheet("06_Data_Dictionary")
ws.sheet_properties.tabColor = "7F7F7F"
set_col_widths(ws, [4, 32, 22, 78])
landscape(ws)
ws["B2"] = "Data dictionary"
ws["B2"].font = title_font
ws["B3"] = "Field definitions so another analyst can inherit the file."
ws["B3"].font = italic_grey
headers = ["Field", "Tab", "Definition"]
for i, h in enumerate(headers):
    cell = ws.cell(5, 2 + i, h)
    cell.font = header_font; cell.fill = header_fill; cell.border = thin

defs = [
    ("As-of / toggle month", ASSUMP, "Month through which the reforecast holds YTD actuals. Later months are forecast. Dashboard INDEX uses this."),
    ("Merit / raise %", ASSUMP, "Company-wide loaded-cost step-up. Applied from the raise-effective month onward."),
    ("Raise effective month", ASSUMP, "First month the merit is in the run-rate (April in the sample)."),
    ("Hire lead time", ASSUMP, "Months from approved req to start. Hires in the plan are booked in the start month."),
    ("Opening FTE", ASSUMP, "1 Jan 2026 heads by department. Seeds January opening of both plan and reforecast."),
    ("Annual loaded cost", ASSUMP, "Fully loaded $000s per FTE per year (salary, benefits, taxes, typical bonus)."),
    ("Annual attrition %", ASSUMP, "Expected voluntary/involuntary turnover. Monthly rate = annual ÷ 12."),
    ("Open reqs", ASSUMP, "Snapshot of approved unfilled requisitions (pipeline), not yet started."),
    ("Hires (starts)", PLAN + " / " + FCST, "Input. Heads that begin that month."),
    ("Modeled attrition exits", PLAN + " / " + FCST, "ROUND(opening FTE × monthly attrition, 1). Driven by the assumption rate."),
    ("Other exits", PLAN + " / " + FCST, "Input. Restructuring, transfers-out, or one-off exits on top of attrition."),
    ("Ending FTE", PLAN + " / " + FCST, "Opening + hires − total exits."),
    ("Raise factor", COST, "1.0 until the merit month, then 1 + raise %."),
    ("Monthly loaded cost", COST, "Average of opening and ending FTE × (annual loaded / 12) × raise factor."),
    ("Plan vs forecast variance", COST, "Forecast $ − plan $. Positive is a cost overrun vs the original plan."),
    ("FY cost bridge", COST, "Opening run-rate + merit on the base + in-year volume/mix/timing plug, tying to plan FY $."),
    ("Month status", FCST, "Actual if month ≤ as-of, else Forecast."),
    ("Units", "All", "FTE in heads (one decimal). Cost in $000s. Rates in %."),
]
for i, (field, tab, definition) in enumerate(defs):
    r = 6 + i
    ws.cell(r, 2, field).font = bold
    ws.cell(r, 2).border = thin
    ws.cell(r, 3, tab).border = thin
    cell = ws.cell(r, 4, definition)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = thin
    ws.row_dimensions[r].height = 36

ws["B26"] = "All sample numbers are fictional. Built for a public GitHub portfolio — no employer data."
ws["B26"].font = small_grey

out = Path(__file__).resolve().parent / "Northline_Headcount_Staffing_Forecast.xlsx"
wb.save(out)
print("Wrote", out)
print("Sheets:", wb.sheetnames)
